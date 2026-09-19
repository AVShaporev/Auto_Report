"""Импорт объектов и оборудования из Excel (этап 1.1 плана после анализа конкурентов).

Поток:
  1. GET  /api/import/objects/template — .xlsx: листы «Объекты», «Оборудование»,
     «Справочники» (текущие значения справочников организации, выпадающие списки).
  2. POST /api/import/objects/preview — разбор файла без записи: что будет
     создано, какие справочники добавятся, ошибки по строкам.
  3. POST /api/import/objects/commit — тот же разбор; при ошибках — 400 с
     предпросмотром, иначе запись.

Запись идёт через обычные service-функции (create_object,
add_equipment_to_object): лимит объектов тарифа, автогенерация плановых
заявок и журнал действий срабатывают так же, как при ручном вводе.

Справочники: регион, тип строения, тип помещения и периодичность должны
существовать (их мало, выбор из списка). Район, населённый пункт, улица,
система, тип оборудования и оборудование создаются, если их нет. Тип улицы и
населённого пункта берётся из сокращения в начале («ул. Садовая», «г. Москва»).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Optional

from fastapi import HTTPException
from loguru import logger
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select

from config import settings
from database.database import new_session
from model.arial import Arial
from model.contract import Contract
from model.equipment import Equipment
from model.locality import Locality
from model.object import Object
from model.objects_equipment import Objects_Equipment
from model.period import Period
from model.region import Region
from model.spec_arial import Spec_Arial
from model.spec_build import Spec_Build
from model.spec_equipment import Spec_Equipment
from model.spec_locality import Spec_Locality
from model.spec_room import Spec_Room
from model.spec_street import Spec_Street
from model.spec_system import Spec_System
from model.street import Street
from model.user import User
from schema.object import ObjectCreate
from schema.objects_equipment import AddEquipmentToObject
from service.activity_log import log_activity

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000

SHEET_OBJECTS = "Объекты"
SHEET_EQUIPMENT = "Оборудование"
SHEET_CATALOGS = "Справочники"

# (ключ, заголовок в шаблоне, обязательное, подсказка)
OBJECT_COLUMNS = [
    ("name", "Объект", True, "Наименование объекта, 2–200 символов. Уникально в пределах договора."),
    ("region", "Регион", True, "Выберите из списка — регион должен быть в справочнике."),
    ("arial", "Район", True, "Если района нет в справочнике — он будет создан."),
    ("locality", "Населённый пункт", True, "С сокращением типа: «г. Москва», «пос. Коммунарка». Новый — будет создан."),
    ("street", "Улица", True, "С сокращением типа: «ул. Садовая», «пр-т Мира». Новая — будет создана."),
    ("build_number", "Дом", False, "Номер дома / строения, до 50 символов."),
    ("spec_build", "Тип строения", True, "Выберите из списка."),
    ("room_number", "Помещение", False, "Номер помещения, до 50 символов."),
    ("spec_room", "Тип помещения", False, "Выберите из списка (необязательно)."),
    ("period", "Периодичность ТО", True, "Выберите из списка — по ней создаются плановые заявки."),
    ("responsible_face", "Ответственное лицо", True, "ФИО ответственного на объекте, 2–100 символов."),
    ("responsible_contact", "Контакт ответственного", True, "Телефон или e-mail, 5–100 символов."),
    ("description", "Описание", False, "Комментарий к объекту."),
    ("requires_signature", "Подпись обязательна", False,
     "да / нет (пусто = нет). «да» — отчёт нельзя отправить на утверждение без подписи "
     "представителя заказчика знаком в мобильном приложении."),
]

_YES = {"да", "yes", "y", "1", "true", "+", "д"}
_NO = {"нет", "no", "n", "0", "false", "-", "н", ""}
EQUIPMENT_COLUMNS = [
    ("object", "Объект", True, "Точно как на листе «Объекты» или как у объекта, уже заведённого в договоре."),
    ("system", "Система", False, "Например «Автоматическая пожарная сигнализация». Новая — будет создана."),
    ("equipment", "Оборудование", True, "Модель / наименование. Новое — будет создано в справочнике."),
    ("spec_equipment", "Тип оборудования", False, "Нужен для нового оборудования. Новый тип — будет создан."),
    ("count", "Кол-во", True, "Целое число ≥ 1."),
    ("inventory_number", "Инвентарный номер", False, "До 50 символов."),
    ("serial_number", "Серийный номер", False, "До 50 символов."),
    ("installation_date", "Дата монтажа", False, "ДД.ММ.ГГГГ."),
]


def _norm(v: Any) -> str:
    """Для сравнения: без лишних пробелов, без регистра, ё = е."""
    return re.sub(r"\s+", " ", str(v or "")).strip().casefold().replace("ё", "е")


def _clean(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


# ============================================================================
# Справочники организации
# ============================================================================

@dataclass
class Catalogs:
    regions: dict[str, Region]
    arials: dict[str, Arial]
    localities: dict[str, Locality]
    streets: dict[str, Street]
    spec_builds: dict[str, Spec_Build]
    spec_rooms: dict[str, Spec_Room]
    periods: dict[str, Period]
    systems: dict[str, Spec_System]
    spec_equipments: dict[str, Spec_Equipment]
    equipments: dict[str, Equipment]
    spec_localities: list[Spec_Locality]
    spec_streets: list[Spec_Street]
    spec_arials: list[Spec_Arial]


async def _load_catalogs(session) -> Catalogs:
    async def all_(model):
        return list((await session.execute(select(model).order_by(model.id))).scalars().all())

    by_name = lambda rows: {_norm(r.name): r for r in rows}  # noqa: E731
    return Catalogs(
        regions=by_name(await all_(Region)),
        arials=by_name(await all_(Arial)),
        localities=by_name(await all_(Locality)),
        streets=by_name(await all_(Street)),
        spec_builds=by_name(await all_(Spec_Build)),
        spec_rooms=by_name(await all_(Spec_Room)),
        periods=by_name(await all_(Period)),
        systems=by_name(await all_(Spec_System)),
        spec_equipments=by_name(await all_(Spec_Equipment)),
        equipments=by_name(await all_(Equipment)),
        spec_localities=await all_(Spec_Locality),
        spec_streets=await all_(Spec_Street),
        spec_arials=await all_(Spec_Arial),
    )


def _split_typed(value: str, specs: list, default_names: tuple[str, ...]):
    """«ул. Садовая» → (spec «ул.», «Садовая»). Без сокращения — тип по умолчанию.

    Сокращение ищется по short_name и по полному имени типа («улица Садовая»).
    Возвращает (spec или None, имя без типа).
    """
    s = value.strip()
    candidates = []
    for sp in specs:
        for label in (getattr(sp, "short_name", None), sp.name):
            if label:
                candidates.append((label.strip(), sp))
    # Длинные сокращения первыми: «пр-т» раньше «пр.»
    for label, sp in sorted(candidates, key=lambda c: -len(c[0])):
        if _norm(s).startswith(_norm(label) + " ") or (label.endswith(".") and _norm(s).startswith(_norm(label))):
            rest = s[len(label):].strip(" .")
            if rest:
                return sp, rest
    default = next((sp for sp in specs if _norm(sp.name) in default_names), specs[0] if specs else None)
    return default, s


# ============================================================================
# Разбор файла
# ============================================================================

@dataclass
class RowResult:
    sheet: str
    row: int
    label: str
    status: str = "new"            # new | exists | skip | error
    errors: list[str] = field(default_factory=list)
    note: Optional[str] = None
    data: dict = field(default_factory=dict)


@dataclass
class ImportPlan:
    contract_id: int
    objects: list[RowResult]
    equipment: list[RowResult]
    new_catalogs: dict[str, list[str]]
    global_errors: list[str]

    @property
    def has_errors(self) -> bool:
        return bool(self.global_errors) or any(
            r.status == "error" for r in self.objects + self.equipment
        )

    def to_dict(self) -> dict:
        def rows(items):
            return [
                {"sheet": r.sheet, "row": r.row, "label": r.label, "status": r.status,
                 "errors": r.errors, "note": r.note}
                for r in items
            ]
        return {
            "contract_id": self.contract_id,
            "summary": {
                "objects_new": sum(r.status == "new" for r in self.objects),
                "objects_existing": sum(r.status == "exists" for r in self.objects),
                "equipment_new": sum(r.status == "new" for r in self.equipment),
                "equipment_skipped": sum(r.status == "skip" for r in self.equipment),
                "errors": sum(r.status == "error" for r in self.objects + self.equipment)
                          + len(self.global_errors),
            },
            "can_import": not self.has_errors and any(
                r.status == "new" for r in self.objects + self.equipment
            ),
            "global_errors": self.global_errors,
            "new_catalogs": self.new_catalogs,
            "objects": rows(self.objects),
            "equipment": rows(self.equipment),
        }


def _read_sheet(wb, title: str, columns) -> tuple[list[tuple[int, dict]], list[str]]:
    """Строки листа как {key: value}. Колонки ищутся по заголовку (порядок любой)."""
    if title not in wb.sheetnames:
        return [], [f"В файле нет листа «{title}». Скачайте шаблон и заполните его."]
    ws = wb[title]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    index = {}
    for i, cell in enumerate(header):
        key = _norm(str(cell or "").replace("*", ""))
        for col_key, caption, _req, _hint in columns:
            if key == _norm(caption):
                index[col_key] = i
    missing = [cap for k, cap, req, _ in columns if req and k not in index]
    if missing:
        return [], [f"Лист «{title}»: нет колонок {', '.join('«' + m + '»' for m in missing)}. "
                     f"Используйте шаблон."]
    out = []
    for n, values in enumerate(rows, start=2):
        rec = {k: (values[i] if i < len(values) else None) for k, i in index.items()}
        if all(_clean(v) is None for v in rec.values()):
            continue
        out.append((n, rec))
        if len(out) > MAX_ROWS:
            return out, [f"Лист «{title}»: больше {MAX_ROWS} строк — разбейте файл на части."]
    return out, []


def _parse_date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(s)


def _len_err(errors, caption, value, lo, hi):
    if value is not None and not (lo <= len(value) <= hi):
        errors.append(f"«{caption}»: от {lo} до {hi} символов (сейчас {len(value)})")


async def build_plan(content: bytes, contract_id: int) -> ImportPlan:
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="Файл больше 5 МБ")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(status_code=400, detail="Не удалось прочитать файл. Нужен .xlsx (Excel 2007+).")

    async with new_session() as session:
        contract = await session.get(Contract, contract_id)
        if not contract:
            raise HTTPException(status_code=404, detail=f"Договор с id {contract_id} не найден")
        cat = await _load_catalogs(session)
        existing_objects = {
            _norm(o.name): o for o in (await session.execute(
                select(Object).where(Object.contract_id == contract_id)
            )).scalars().all()
        }
        existing_links = {
            (oe.object_id, oe.equipment_id) for oe in (await session.execute(
                select(Objects_Equipment).where(
                    Objects_Equipment.object_id.in_([o.id for o in existing_objects.values()] or [0])
                )
            )).scalars().all()
        }
        objects_total = (await session.execute(select(func.count(Object.id)))).scalar_one()

    obj_rows, err1 = _read_sheet(wb, SHEET_OBJECTS, OBJECT_COLUMNS)
    eq_rows, err2 = _read_sheet(wb, SHEET_EQUIPMENT, EQUIPMENT_COLUMNS)
    global_errors = err1 + err2
    new_cat: dict[str, dict[str, str]] = {k: {} for k in (
        "arials", "localities", "streets", "systems", "spec_equipments", "equipments")}

    # ---------- Объекты ----------
    objects: list[RowResult] = []
    seen: dict[str, int] = {}
    for n, rec in obj_rows:
        v = {k: _clean(rec.get(k)) for k, *_ in OBJECT_COLUMNS}
        r = RowResult(SHEET_OBJECTS, n, v["name"] or "(без названия)")
        for key, caption, required, _ in OBJECT_COLUMNS:
            if required and not v[key]:
                r.errors.append(f"не заполнено «{caption}»")
        _len_err(r.errors, "Объект", v["name"], 2, 200)
        _len_err(r.errors, "Ответственное лицо", v["responsible_face"], 2, 100)
        _len_err(r.errors, "Контакт ответственного", v["responsible_contact"], 5, 100)
        _len_err(r.errors, "Дом", v["build_number"], 1, 50)
        _len_err(r.errors, "Помещение", v["room_number"], 1, 50)
        flag = _norm(v.get("requires_signature") or "")
        if flag not in _YES and flag not in _NO:
            r.errors.append(f"«Подпись обязательна»: «{v['requires_signature']}» — нужно «да» или «нет»")
        v["requires_signature_bool"] = flag in _YES

        key = _norm(v["name"])
        if v["name"] and key in seen:
            r.errors.append(f"объект с таким названием уже есть в строке {seen[key]}")
        elif v["name"]:
            seen[key] = n

        def must(value, catalog, caption):
            if value and _norm(value) not in catalog:
                opts = ", ".join(sorted(x.name for x in catalog.values())[:12])
                r.errors.append(f"«{caption}»: «{value}» нет в справочнике (есть: {opts or 'пусто'})")
        must(v["region"], cat.regions, "Регион")
        must(v["spec_build"], cat.spec_builds, "Тип строения")
        must(v["spec_room"], cat.spec_rooms, "Тип помещения")
        must(v["period"], cat.periods, "Периодичность ТО")

        # Создаваемые справочники: район / НП / улица
        if v["arial"] and _norm(v["arial"]) not in cat.arials:
            new_cat["arials"].setdefault(_norm(v["arial"]), v["arial"])
        if v["locality"]:
            _sp, lname = _split_typed(v["locality"], cat.spec_localities, ("город",))
            if _norm(lname) not in cat.localities:
                new_cat["localities"].setdefault(_norm(lname), v["locality"])
        if v["street"]:
            _sp, sname = _split_typed(v["street"], cat.spec_streets, ("улица",))
            if _norm(sname) not in cat.streets:
                new_cat["streets"].setdefault(_norm(sname), v["street"])

        if r.errors:
            r.status = "error"
        elif key in existing_objects:
            r.status = "exists"
            r.note = "уже есть в договоре — не создаётся"
        r.data = v
        objects.append(r)

    new_objects = {_norm(r.data["name"]) for r in objects if r.status == "new"}
    if settings.MAX_OBJECTS is not None and objects_total + len(new_objects) > settings.MAX_OBJECTS:
        global_errors.append(
            f"Лимит тарифа — {settings.MAX_OBJECTS} объектов: сейчас {objects_total}, "
            f"в файле новых {len(new_objects)}."
        )

    # ---------- Оборудование ----------
    equipment: list[RowResult] = []
    planned_links: dict[tuple[str, str], int] = {}
    for n, rec in eq_rows:
        v = {k: _clean(rec.get(k)) for k, *_ in EQUIPMENT_COLUMNS}
        r = RowResult(SHEET_EQUIPMENT, n, f"{v['equipment'] or '?'} → {v['object'] or '?'}")
        for key, caption, required, _ in EQUIPMENT_COLUMNS:
            if required and v[key] in (None, ""):
                r.errors.append(f"не заполнено «{caption}»")
        okey = _norm(v["object"])
        if v["object"] and okey not in new_objects and okey not in existing_objects:
            in_file = okey in seen
            r.errors.append(
                f"объект «{v['object']}» " + ("с ошибками на листе «Объекты»" if in_file
                                               else "не найден ни в файле, ни в договоре")
            )
        count = None
        if v["count"] is not None:
            try:
                count = int(float(str(v["count"]).replace(",", ".")))
                if count < 1 or float(str(v["count"]).replace(",", ".")) != count:
                    raise ValueError
            except ValueError:
                r.errors.append(f"«Кол-во»: нужно целое число ≥ 1 (сейчас «{v['count']}»)")
        try:
            inst = _parse_date(rec.get("installation_date"))
        except ValueError:
            inst = None
            r.errors.append(f"«Дата монтажа»: не дата (нужно ДД.ММ.ГГГГ)")
        _len_err(r.errors, "Оборудование", v["equipment"], 1, 200)
        _len_err(r.errors, "Инвентарный номер", v["inventory_number"], 1, 50)
        _len_err(r.errors, "Серийный номер", v["serial_number"], 1, 50)

        ekey = _norm(v["equipment"])
        if v["equipment"] and ekey not in cat.equipments:
            if not v["spec_equipment"]:
                r.errors.append("новое оборудование — заполните «Тип оборудования»")
            else:
                new_cat["equipments"].setdefault(ekey, v["equipment"])
                if _norm(v["spec_equipment"]) not in cat.spec_equipments:
                    new_cat["spec_equipments"].setdefault(_norm(v["spec_equipment"]), v["spec_equipment"])
        if v["system"] and _norm(v["system"]) not in cat.systems:
            new_cat["systems"].setdefault(_norm(v["system"]), v["system"])

        if not r.errors and v["object"] and v["equipment"]:
            pair = (okey, ekey)
            if pair in planned_links:
                r.errors.append(f"это оборудование на этом объекте уже в строке {planned_links[pair]} — "
                                f"объедините количество")
            else:
                planned_links[pair] = n
            existing_obj = existing_objects.get(okey)
            existing_eq = cat.equipments.get(ekey)
            if (not r.errors and existing_obj and existing_eq
                    and (existing_obj.id, existing_eq.id) in existing_links):
                r.status = "skip"
                r.note = "уже есть на объекте — пропущено"
        if r.errors:
            r.status = "error"
        v["count_int"], v["installation"] = count, inst
        r.data = v
        equipment.append(r)

    return ImportPlan(
        contract_id=contract_id,
        objects=objects,
        equipment=equipment,
        new_catalogs={k: sorted(d.values()) for k, d in new_cat.items()},
        global_errors=global_errors,
    )


# ============================================================================
# Запись
# ============================================================================

async def commit_plan(content: bytes, contract_id: int, current_user: User) -> dict:
    """Импорт. Ошибки в файле → 400 с предпросмотром (ничего не записано)."""
    # Импортные права = ручные: создавать объекты и добавлять на них оборудование.
    from service.object import check_permission as obj_perm, create_object
    from service.objects_equipment import add_equipment_to_object

    await obj_perm(current_user, "object_create", "импорта объектов")
    await obj_perm(current_user, "object_equipment_create", "импорта оборудования")

    plan = await build_plan(content, contract_id)
    if plan.has_errors:
        raise HTTPException(status_code=400, detail={"message": "В файле есть ошибки", "preview": plan.to_dict()})

    # 1. Недостающие справочники — одной транзакцией.
    async with new_session() as session:
        cat = await _load_catalogs(session)
        default_arial = cat.spec_arials[0] if cat.spec_arials else None

        for r in plan.objects:
            if r.status != "new":
                continue
            v = r.data
            if _norm(v["arial"]) not in cat.arials:
                if default_arial is None:
                    raise HTTPException(status_code=400, detail="В справочнике нет типов районов — добавьте хотя бы один")
                a = Arial(name=v["arial"], spec_arial_id=default_arial.id)
                session.add(a)
                cat.arials[_norm(v["arial"])] = a
            sp, lname = _split_typed(v["locality"], cat.spec_localities, ("город",))
            if _norm(lname) not in cat.localities:
                loc = Locality(name=lname, spec_locality_id=sp.id if sp else None)
                session.add(loc)
                cat.localities[_norm(lname)] = loc
            sp, sname = _split_typed(v["street"], cat.spec_streets, ("улица",))
            if _norm(sname) not in cat.streets:
                if sp is None:
                    raise HTTPException(status_code=400, detail="В справочнике нет типов улиц — добавьте хотя бы один")
                st = Street(name=sname, spec_street_id=sp.id)
                session.add(st)
                cat.streets[_norm(sname)] = st
        await session.flush()

        for r in plan.equipment:
            if r.status != "new":
                continue
            v = r.data
            if v["system"] and _norm(v["system"]) not in cat.systems:
                s = Spec_System(name=v["system"])
                session.add(s)
                cat.systems[_norm(v["system"])] = s
            if _norm(v["equipment"]) not in cat.equipments:
                if _norm(v["spec_equipment"]) not in cat.spec_equipments:
                    se = Spec_Equipment(name=v["spec_equipment"])
                    session.add(se)
                    cat.spec_equipments[_norm(v["spec_equipment"])] = se
                await session.flush()
                sys_ = cat.systems.get(_norm(v["system"])) if v["system"] else None
                eq = Equipment(
                    name=v["equipment"],
                    spec_equipment_id=cat.spec_equipments[_norm(v["spec_equipment"])].id,
                    spec_system_id=sys_.id if sys_ else None,
                )
                session.add(eq)
                cat.equipments[_norm(v["equipment"])] = eq
        await session.flush()
        ids = {
            "region": {k: x.id for k, x in cat.regions.items()},
            "arial": {k: x.id for k, x in cat.arials.items()},
            "locality": {k: x.id for k, x in cat.localities.items()},
            "street": {k: x.id for k, x in cat.streets.items()},
            "spec_build": {k: x.id for k, x in cat.spec_builds.items()},
            "spec_room": {k: x.id for k, x in cat.spec_rooms.items()},
            "period": {k: x.id for k, x in cat.periods.items()},
            "system": {k: x.id for k, x in cat.systems.items()},
            "equipment": {k: x.id for k, x in cat.equipments.items()},
        }
        await session.commit()
        existing = {
            _norm(o.name): o.id for o in (await session.execute(
                select(Object).where(Object.contract_id == contract_id)
            )).scalars().all()
        }

    # 2. Объекты — обычным create_object (лимит, автогенерация заявок, журнал).
    created_objects, failed = 0, []
    object_ids = dict(existing)
    for r in plan.objects:
        if r.status != "new":
            continue
        v = r.data
        try:
            obj = await create_object(ObjectCreate(
                name=v["name"],
                build_number=v["build_number"],
                room_number=v["room_number"],
                responsible_face=v["responsible_face"],
                responsible_faces_contact=v["responsible_contact"],
                region_id=ids["region"][_norm(v["region"])],
                arial_id=ids["arial"][_norm(v["arial"])],
                locality_id=ids["locality"][_norm(_split_typed(v["locality"], cat.spec_localities, ("город",))[1])],
                street_id=ids["street"][_norm(_split_typed(v["street"], cat.spec_streets, ("улица",))[1])],
                spec_build_id=ids["spec_build"][_norm(v["spec_build"])],
                spec_room_id=ids["spec_room"][_norm(v["spec_room"])] if v["spec_room"] else None,
                period_id=ids["period"][_norm(v["period"])],
                contract_id=contract_id,
                description=v["description"],
                requires_signature=v.get("requires_signature_bool", False),
            ), current_user)
            object_ids[_norm(v["name"])] = obj.id
            created_objects += 1
        except HTTPException as e:
            failed.append({"sheet": r.sheet, "row": r.row, "label": r.label, "error": str(e.detail)})
        except Exception as e:  # noqa: BLE001
            logger.exception("import: объект строки {} не создан", r.row)
            failed.append({"sheet": r.sheet, "row": r.row, "label": r.label, "error": type(e).__name__})

    # 3. Оборудование на объекты.
    created_links = 0
    for r in plan.equipment:
        if r.status != "new":
            continue
        v = r.data
        obj_id = object_ids.get(_norm(v["object"]))
        if not obj_id:
            failed.append({"sheet": r.sheet, "row": r.row, "label": r.label, "error": "объект не создан"})
            continue
        try:
            await add_equipment_to_object(obj_id, AddEquipmentToObject(
                equipment_id=ids["equipment"][_norm(v["equipment"])],
                count=v["count_int"],
                inventory_number=v["inventory_number"],
                serial_number=v["serial_number"],
                installation_date=v["installation"],
                spec_system_id=ids["system"][_norm(v["system"])] if v["system"] else None,
            ), current_user)
            created_links += 1
        except HTTPException as e:
            failed.append({"sheet": r.sheet, "row": r.row, "label": r.label, "error": str(e.detail)})

    async with new_session() as session:
        await log_activity(
            session, current_user, action="create", entity="object", entity_id=None,
            summary=f"Импорт из Excel: {created_objects} объектов, {created_links} позиций оборудования",
            details={"contract_id": contract_id, "new_catalogs": plan.new_catalogs, "failed": failed},
        )
        await session.commit()

    return {
        "objects_created": created_objects,
        "equipment_created": created_links,
        "new_catalogs": plan.new_catalogs,
        "failed": failed,
    }


# ============================================================================
# Шаблон
# ============================================================================

async def build_template() -> bytes:
    async with new_session() as session:
        cat = await _load_catalogs(session)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2B6CD9")
    req_fill = PatternFill("solid", fgColor="1E4FA6")

    # Справочники — на отдельном листе, из них выпадающие списки.
    ws_c = wb.active
    ws_c.title = SHEET_CATALOGS
    lists = [
        ("Регионы", cat.regions), ("Районы", cat.arials), ("Населённые пункты", cat.localities),
        ("Улицы", cat.streets), ("Типы строений", cat.spec_builds), ("Типы помещений", cat.spec_rooms),
        ("Периодичность ТО", cat.periods), ("Системы", cat.systems),
        ("Типы оборудования", cat.spec_equipments), ("Оборудование", cat.equipments),
    ]
    ranges = {}
    for col, (caption, catalog) in enumerate(lists, start=1):
        letter = get_column_letter(col)
        ws_c.cell(row=1, column=col, value=caption).font = Font(bold=True)
        names = sorted({x.name for x in catalog.values()}, key=str.casefold)
        for i, name in enumerate(names, start=2):
            ws_c.cell(row=i, column=col, value=name)
        ws_c.column_dimensions[letter].width = 34
        ranges[caption] = f"'{SHEET_CATALOGS}'!${letter}$2:${letter}${max(len(names) + 1, 2)}"

    def sheet(title, columns, lookups):
        ws = wb.create_sheet(title, index=len(wb.sheetnames) - 1)
        for col, (key, caption, required, hint) in enumerate(columns, start=1):
            c = ws.cell(row=1, column=col, value=caption + (" *" if required else ""))
            c.font = head_font
            c.fill = req_fill if required else head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.comment = Comment(hint, "Cool-Doc")
            ws.column_dimensions[get_column_letter(col)].width = max(16, len(caption) + 6)
            if key in lookups:
                strict, catalog_caption = lookups[key]
                formula = '"да,нет"' if catalog_caption == "__yes_no__" else ranges[catalog_caption]
                dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                                    showErrorMessage=strict)
                if strict:
                    dv.error = "Выберите значение из списка (лист «Справочники»)."
                    dv.errorTitle = caption
                ws.add_data_validation(dv)
                dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{MAX_ROWS + 1}")
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        return ws

    ws_o = sheet(SHEET_OBJECTS, OBJECT_COLUMNS, {
        "requires_signature": (True, "__yes_no__"),
        "region": (True, "Регионы"), "arial": (False, "Районы"), "locality": (False, "Населённые пункты"),
        "spec_build": (True, "Типы строений"), "spec_room": (True, "Типы помещений"),
        "period": (True, "Периодичность ТО"),
    })
    ws_e = sheet(SHEET_EQUIPMENT, EQUIPMENT_COLUMNS, {
        "system": (False, "Системы"), "equipment": (False, "Оборудование"),
        "spec_equipment": (False, "Типы оборудования"),
    })
    ws_o.column_dimensions["A"].width = 34
    ws_e.column_dimensions["A"].width = 34
    ws_e.column_dimensions["C"].width = 40
    wb.active = wb.sheetnames.index(SHEET_OBJECTS)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
