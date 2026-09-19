"""Импорт объектов и оборудования из Excel (service/import_objects.py)."""
import io

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from model.equipment import Equipment
from model.locality import Locality
from model.object import Object
from model.objects_equipment import Objects_Equipment
from model.street import Street

OBJ_HEAD = ["Объект *", "Регион *", "Район *", "Населённый пункт *", "Улица *", "Дом",
            "Тип строения *", "Помещение", "Тип помещения", "Периодичность ТО *",
            "Ответственное лицо *", "Контакт ответственного *", "Описание"]
EQ_HEAD = ["Объект *", "Система", "Оборудование *", "Тип оборудования", "Кол-во *",
           "Инвентарный номер", "Серийный номер", "Дата монтажа"]


def _xlsx(objects: list[list], equipment: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Объекты"
    ws.append(OBJ_HEAD)
    for r in objects:
        ws.append(r)
    ws2 = wb.create_sheet("Оборудование")
    ws2.append(EQ_HEAD)
    for r in equipment:
        ws2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _obj_row(name="БЦ Северный", region="Тестовый регион", **kw):
    row = {"name": name, "region": region, "arial": "Новый район", "locality": "г. Новгород",
           "street": "ул. Новая", "house": "5", "build": "Здание", "room": None, "room_type": None,
           "period": "Месяц", "face": "Петров П.П.", "contact": "+7 900 000-00-00", "descr": None}
    row.update(kw)
    return list(row.values())


VALID = lambda: _xlsx(  # noqa: E731
    [_obj_row()],
    [["БЦ Северный", None, "Огнетушитель ОП-5", None, 3, "ИНВ-1", None, "01.02.2026"],
     ["БЦ Северный", "Автоматическая пожарная сигнализация", "Извещатель ИП 212-45", "Извещатель", 64, None, None, None]],
)


async def _post(client, headers, url, content, contract_id):
    return await client.post(
        url, headers=headers, data={"contract_id": str(contract_id)},
        files={"file": ("import.xlsx", content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


async def test_template_has_sheets_and_dropdowns(client: AsyncClient, db_session, superadmin_token,
                                                 superadmin_user, auth_headers, reference_data):
    resp = await client.get("/api/import/objects/template", headers=auth_headers(superadmin_token))
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Объекты", "Оборудование", "Справочники"]
    assert wb["Объекты"]["B1"].value == "Регион *"
    assert "Тестовый регион" in [c.value for c in wb["Справочники"]["A"]]
    assert len(wb["Объекты"].data_validations.dataValidation) >= 4


async def test_preview_then_commit_creates_objects_equipment_and_catalogs(
    client: AsyncClient, db_session, superadmin_token, superadmin_user, auth_headers, reference_data,
):
    h = auth_headers(superadmin_token)
    contract_id = reference_data["contract"].id
    # id до expire_all — иначе ленивая подгрузка вне greenlet
    spec_locality_id = reference_data["spec_locality"].id
    spec_street_id = reference_data["spec_street"].id

    resp = await _post(client, h, "/api/import/objects/preview", VALID(), contract_id)
    assert resp.status_code == 200, resp.text
    p = resp.json()
    assert p["can_import"] is True, p
    assert p["summary"] == {"objects_new": 1, "objects_existing": 0, "equipment_new": 2,
                            "equipment_skipped": 0, "errors": 0}
    assert p["new_catalogs"]["arials"] == ["Новый район"]
    assert p["new_catalogs"]["localities"] == ["г. Новгород"]
    assert p["new_catalogs"]["streets"] == ["ул. Новая"]
    assert p["new_catalogs"]["equipments"] == ["Извещатель ИП 212-45"]
    assert p["new_catalogs"]["spec_equipments"] == ["Извещатель"]
    assert p["new_catalogs"]["systems"] == ["Автоматическая пожарная сигнализация"]

    resp = await _post(client, h, "/api/import/objects/commit", VALID(), contract_id)
    assert resp.status_code == 200, resp.text
    r = resp.json()
    assert (r["objects_created"], r["equipment_created"], r["failed"]) == (1, 2, [])

    db_session.expire_all()
    obj = (await db_session.execute(select(Object).where(Object.name == "БЦ Северный"))).scalar_one()
    loc = await db_session.get(Locality, obj.locality_id)
    street = await db_session.get(Street, obj.street_id)
    # Тип отделён от названия: «г. Новгород» → «Новгород» с типом «г.»
    assert (loc.name, loc.spec_locality_id) == ("Новгород", spec_locality_id)
    assert (street.name, street.spec_street_id) == ("Новая", spec_street_id)
    assert obj.number_in_contract == 2  # в reference_data уже есть объект №1
    links = (await db_session.execute(
        select(Objects_Equipment).where(Objects_Equipment.object_id == obj.id)
    )).scalars().all()
    assert sorted(l.count for l in links) == [3, 64]
    new_eq = (await db_session.execute(
        select(Equipment).where(Equipment.name == "Извещатель ИП 212-45")
    )).scalar_one()
    assert new_eq.spec_system_id is not None

    # Повторная загрузка того же файла ничего не дублирует
    resp = await _post(client, h, "/api/import/objects/preview", VALID(), contract_id)
    p = resp.json()
    assert p["summary"]["objects_existing"] == 1
    assert p["summary"]["equipment_skipped"] == 2
    assert p["can_import"] is False


async def test_errors_block_import(client: AsyncClient, db_session, superadmin_token,
                                   superadmin_user, auth_headers, reference_data):
    h = auth_headers(superadmin_token)
    content = _xlsx(
        [_obj_row(region="Марс"),                      # нет в справочнике
         _obj_row(name="Склад", face=None),            # не заполнено
         _obj_row(name="Склад")],                      # дубль названия
        [["Склад", None, "Огнетушитель ОП-5", None, "два", None, None, None],   # не число
         ["Нет такого", None, "Огнетушитель ОП-5", None, 1, None, None, None],  # нет объекта
         ["Склад", None, "Совсем новое", None, 1, None, None, None]],           # нет типа
    )
    contract_id = reference_data["contract"].id
    resp = await _post(client, h, "/api/import/objects/preview", content, contract_id)
    p = resp.json()
    assert p["can_import"] is False
    errs = {(r["sheet"], r["row"]): " | ".join(r["errors"]) for r in p["objects"] + p["equipment"] if r["errors"]}
    assert "Марс" in errs[("Объекты", 2)]
    assert "Ответственное лицо" in errs[("Объекты", 3)]
    assert "строке 3" in errs[("Объекты", 4)]
    assert "целое число" in errs[("Оборудование", 2)]
    assert "не найден" in errs[("Оборудование", 3)]
    assert "Тип оборудования" in errs[("Оборудование", 4)]

    resp = await _post(client, h, "/api/import/objects/commit", content, contract_id)
    assert resp.status_code == 400
    assert resp.json()["detail"]["preview"]["can_import"] is False
    db_session.expire_all()
    assert (await db_session.execute(select(Object).where(Object.name == "Склад"))).first() is None


async def test_bad_file_and_permissions(client: AsyncClient, db_session, regular_token, admin_token,
                                        superadmin_token, auth_headers, reference_data):
    resp = await _post(client, auth_headers(superadmin_token), "/api/import/objects/preview",
                       b"not an excel", reference_data["contract"].id)
    assert resp.status_code == 400 and ".xlsx" in resp.json()["detail"]
    resp = await client.get("/api/import/objects/template", headers=auth_headers(regular_token))
    assert resp.status_code == 403
