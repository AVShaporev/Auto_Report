"""Метки docxtpl-шаблонов (1.0.67): report.*, order.period_month/type/…,
works / works_groups, issues_open / issues_fixed в акте заявки;
reports / issues / equipment_groups / works в журнале объекта.

Проверяем и сам словарь контекста, и настоящий рендер .docx по шаблону с
циклами {%tr %} — так, как это сделает пользовательский шаблон.
"""
import io
from datetime import date

from docx import Document
from docxtpl import DocxTemplate

from model.operation import Operation
from model.order import Order
from model.report import Report
from model.spec_order import Spec_Order
from model.spec_order_status import Spec_Order_Status
from model.spec_report_status import Spec_Report_Status
from service import render_docx
from tests.factories import create_issue


async def _seed(db_session, reference_data, user) -> dict:
    ref = reference_data
    plan = Spec_Order(name="Плановое ТО", short_name="ТО", code="plan", sla_kind="manual")
    order_status = Spec_Order_Status(name="Выполнена", is_default=True)
    approved = Spec_Report_Status(name="Утверждён", is_default=True)
    draft = Spec_Report_Status(name="В работе")
    db_session.add_all([plan, order_status, approved, draft])
    await db_session.commit()

    op_check = Operation(name="Внешний осмотр", period_id=ref["period"].id)
    op_check.spec_equipments = [ref["spec_equipment"]]
    op_weigh = Operation(name="Контроль массы заряда")
    op_weigh.spec_equipments = [ref["spec_equipment"]]
    db_session.add_all([op_check, op_weigh])
    await db_session.commit()

    report = Report(
        number="ОТЧ-9", status_id=approved.id, period_id=ref["period"].id,
        contract_id=ref["contract"].id, object_id=ref["object"].id, user_id=user.id,
        created_at=date(2026, 9, 18), description="Проведён осмотр, замечаний нет",
    )
    old_draft = Report(
        number="ОТЧ-8", status_id=draft.id, period_id=ref["period"].id,
        contract_id=ref["contract"].id, object_id=ref["object"].id, user_id=user.id,
        created_at=date(2026, 8, 18), description="Черновик — в журнал не попадает",
    )
    db_session.add_all([report, old_draft])
    await db_session.commit()

    order = Order(
        number="1/09/2026/ТО/1", spec_order_id=plan.id, contract_id=ref["contract"].id,
        object_id=ref["object"].id, user_id=user.id, assigned_to_id=user.id,
        status_id=order_status.id, report_id=report.id,
        period_start_date=date(2026, 9, 1), due_date=date(2026, 9, 30),
        created_at=date(2026, 9, 1), description="Ежемесячное ТО",
    )
    db_session.add(order)
    await db_session.commit()

    open_issue = await create_issue(
        db_session, object_equipment_id=ref["objects_equipment"].id,
        priority_id=ref["spec_priority"]["high"].id, status_id=ref["spec_status"]["new"].id,
        reported_by_id=user.id, title="Сорвана пломба", detected_date=date(2026, 9, 18),
        number="Н-1",
    )
    fixed_issue = await create_issue(
        db_session, object_equipment_id=ref["objects_equipment"].id,
        priority_id=ref["spec_priority"]["medium"].id, status_id=ref["spec_status"]["resolved"].id,
        reported_by_id=user.id, title="Потёк корпус", detected_date=date(2026, 8, 2),
        number="Н-2",
    )
    # Устранена заявкой на устранение — этих полей нет в фабрике.
    fixed_issue.is_resolved = True
    fixed_issue.resolved_date = date(2026, 9, 5)
    fixed_issue.order_id = order.id
    await db_session.commit()
    return {"order_id": order.id, "open_issue": open_issue, "fixed_issue": fixed_issue}


async def test_order_context_has_report_works_and_issues(client, db_session, superadmin_user, reference_data):
    user = superadmin_user["user"]
    seeded = await _seed(db_session, reference_data, user)

    order = await render_docx._load_order_for_docx(db_session, seeded["order_id"])
    ctx = render_docx._build_context(order)

    assert ctx["order"]["type"] == "Плановое ТО"
    assert ctx["order"]["status"] == "Выполнена"
    assert ctx["order"]["period_month"] == "сентябрь 2026"
    assert ctx["order"]["period_start"] == "01.09.2026"
    assert ctx["order"]["due_date"] == "30.09.2026"
    assert ctx["order"]["created_at_long"] == "«01» сентября 2026 г."
    assert ctx["order"]["assigned_to"] == (user.full_name or "")

    assert ctx["report"]["number"] == "ОТЧ-9"
    assert ctx["report"]["date"] == "18.09.2026"
    assert ctx["report"]["date_long"] == "«18» сентября 2026 г."
    assert ctx["report"]["description"] == "Проведён осмотр, замечаний нет"
    assert ctx["report"]["is_approved"] is True

    assert [w["name"] for w in ctx["works"]] == ["Внешний осмотр", "Контроль массы заряда"]
    assert ctx["works"][0]["period"] == "Месяц"
    assert ctx["works_groups"][0]["equipment_type"] == "Огнетушитель"
    assert [r["index"] for r in ctx["works_groups"][0]["rows"]] == ["1.1", "1.2"]

    assert [i["title"] for i in ctx["issues_open"]] == ["Сорвана пломба"]
    assert ctx["issues_open"][0]["equipment"] == "Огнетушитель ОП-5"
    assert ctx["issues_open"][0]["priority"] == "Высокий"
    assert [i["title"] for i in ctx["issues_fixed"]] == ["Потёк корпус"]
    assert ctx["issues_fixed"][0]["resolved_date"] == "05.09.2026"


async def test_order_context_without_report_is_empty_strings(db_session, superadmin_user, reference_data):
    """Шаблон с {{ report.* }} не падает, если отчёта ещё нет."""
    ref, user = reference_data, superadmin_user["user"]
    plan = Spec_Order(name="Аварийная", short_name="АВР", code="emergency", sla_kind="manual")
    st = Spec_Order_Status(name="Новая", is_default=True)
    db_session.add_all([plan, st])
    await db_session.commit()
    order = Order(number="A-1", spec_order_id=plan.id, contract_id=ref["contract"].id,
                  object_id=ref["object"].id, user_id=user.id, status_id=st.id,
                  created_at=date(2026, 3, 10))
    db_session.add(order)
    await db_session.commit()

    ctx = render_docx._build_context(await render_docx._load_order_for_docx(db_session, order.id))
    assert ctx["report"]["description"] == "" and ctx["report"]["is_approved"] is False
    assert ctx["order"]["period_month"] == "март 2026"  # нет периода → месяц создания
    assert ctx["order"]["assigned_to"] == ""
    assert ctx["issues_fixed"] == []


async def test_journal_context_lists_approved_reports_and_issues(db_session, superadmin_user, reference_data):
    user = superadmin_user["user"]
    await _seed(db_session, reference_data, user)

    obj = await render_docx._load_object_for_journal(db_session, reference_data["object"].id)
    ctx = render_docx._build_journal_context(obj)

    # Только утверждённые отчёты, по дате
    assert [r["number"] for r in ctx["reports"]] == ["ОТЧ-9"]
    assert ctx["reports"][0]["order_type"] == "Плановое ТО"
    assert ctx["reports"][0]["order_number"] == "1/09/2026/ТО/1"
    # Все неисправности по дате обнаружения + только открытые
    assert [i["title"] for i in ctx["issues"]] == ["Потёк корпус", "Сорвана пломба"]
    assert [i["title"] for i in ctx["issues_open"]] == ["Сорвана пломба"]
    assert ctx["equipment_groups"][0]["rows"][0]["name"] == "Огнетушитель ОП-5"
    assert len(ctx["works"]) == 2


def _render(template_rows: list[list[str]], context: dict) -> list[list[str]]:
    """Собрать .docx с одной таблицей из строк-шаблонов, отрендерить, вернуть ячейки."""
    doc = Document()
    table = doc.add_table(rows=len(template_rows), cols=len(template_rows[0]))
    for r, row in enumerate(template_rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    tpl = DocxTemplate(buf)
    tpl.render(context)
    out = io.BytesIO()
    tpl.save(out)
    out.seek(0)
    rendered = Document(out).tables[0]
    return [[cell.text for cell in row.cells] for row in rendered.rows]


async def test_journal_reports_table_renders_rows(db_session, superadmin_user, reference_data):
    """Таблица журнала {%tr for r in reports %} разворачивается в строки."""
    user = superadmin_user["user"]
    await _seed(db_session, reference_data, user)
    obj = await render_docx._load_object_for_journal(db_session, reference_data["object"].id)
    ctx = render_docx._build_journal_context(obj)

    rows = _render([
        ["Дата", "Работы", "Исполнитель"],
        ["{%tr for r in reports %}", "", ""],
        ["{{ r.date }}", "{{ r.description }}", "{{ r.engineer }}"],
        ["{%tr endfor %}", "", ""],
    ], ctx)
    assert rows[0] == ["Дата", "Работы", "Исполнитель"]
    assert rows[1][:2] == ["18.09.2026", "Проведён осмотр, замечаний нет"]
    assert len(rows) == 2
