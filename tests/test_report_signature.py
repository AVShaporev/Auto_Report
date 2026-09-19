"""Подпись заказчика под отчётом (1.0.69): приём с отчётом, замена, удаление,
запрет после утверждения, картинка по GET, метки в акте."""
import base64
import io
from datetime import date

from httpx import AsyncClient
from PIL import Image, ImageDraw

from model.order import Order
from model.report import Report
from model.spec_order import Spec_Order
from model.spec_order_status import Spec_Order_Status
from model.spec_report_status import Spec_Report_Status
from config import MEDIA_PATH
from docx import Document
from service import render_docx


def _signature_png(transparent=True) -> str:
    img = Image.new("RGBA", (600, 240), (0, 0, 0, 0) if transparent else (255, 255, 255, 255))
    d = ImageDraw.Draw(img)
    d.line([(60, 170), (180, 60), (300, 180), (420, 70), (540, 150)], fill=(20, 30, 90, 255), width=6)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


SIG = lambda **kw: {"image": _signature_png(), "signer_name": "Орлова Мария Сергеевна",  # noqa: E731
                    "signer_position": "Главный инженер", **kw}


async def _seed(db_session, ref, user):
    plan = Spec_Order(name="Плановое ТО", short_name="ТО", code="plan", sla_kind="manual")
    st = Spec_Order_Status(name="Новая", is_default=True)
    work = Spec_Report_Status(name="В работе", is_default=True)
    approved = Spec_Report_Status(name="Утверждён")
    db_session.add_all([plan, st, work, approved])
    await db_session.commit()
    order = Order(number="1/09/2026/ТО/1", spec_order_id=plan.id, contract_id=ref["contract"].id,
                  object_id=ref["object"].id, user_id=user.id, status_id=st.id,
                  created_at=date(2026, 9, 1))
    db_session.add(order)
    await db_session.commit()
    return {"order_id": order.id, "approved_id": approved.id}


async def test_report_created_with_signature_then_replaced_and_cleared(
    client: AsyncClient, db_session, superadmin_token, superadmin_user, auth_headers, reference_data,
):
    h = auth_headers(superadmin_token)
    seeded = await _seed(db_session, reference_data, superadmin_user["user"])

    resp = await client.post("/api/report/create", headers=h, json={
        "order_id": seeded["order_id"], "report_period": "2026-09",
        "description": "ТО выполнено", "customer_signature": SIG(signed_at="2026-09-18T11:32:00+03:00"),
    })
    assert resp.status_code == 200, resp.text
    rep = resp.json()
    assert rep["has_signature"] is True
    assert rep["signer_name"] == "Орлова Мария Сергеевна"
    assert rep["signer_position"] == "Главный инженер"
    assert rep["signed_at"].startswith("2026-09-18T08:32")  # в UTC

    resp = await client.get(f"/api/report/{rep['id']}/signature", headers=h)
    assert resp.status_code == 200 and resp.headers["content-type"] == "image/png"
    img = Image.open(io.BytesIO(resp.content))
    # Фон белый (не прозрачный), поля вокруг росчерка обрезаны
    assert img.mode == "L" and img.width < 600 and img.getpixel((0, 0)) == 255

    # Замена подписи
    resp = await client.put(f"/api/report/{rep['id']}", headers=h,
                            json={"customer_signature": SIG(signer_name="Петров П.П.", signer_position=None)})
    assert resp.status_code == 200, resp.text
    assert resp.json()["signer_name"] == "Петров П.П." and resp.json()["signer_position"] is None

    # Правка описания без подписи подпись не трогает
    resp = await client.put(f"/api/report/{rep['id']}", headers=h, json={"description": "Исправил"})
    assert resp.json()["has_signature"] is True

    # null — убрать
    resp = await client.put(f"/api/report/{rep['id']}", headers=h, json={"customer_signature": None})
    assert resp.status_code == 200 and resp.json()["has_signature"] is False
    resp = await client.get(f"/api/report/{rep['id']}/signature", headers=h)
    assert resp.status_code == 404


async def test_signature_validation_and_locked_after_approval(
    client: AsyncClient, db_session, superadmin_token, superadmin_user, auth_headers, reference_data,
):
    h = auth_headers(superadmin_token)
    seeded = await _seed(db_session, reference_data, superadmin_user["user"])
    resp = await client.post("/api/report/create", headers=h, json={
        "order_id": seeded["order_id"], "report_period": "2026-09"})
    rid = resp.json()["id"]

    blank = Image.new("RGBA", (300, 100), (0, 0, 0, 0))
    buf = io.BytesIO()
    blank.save(buf, "PNG")
    empty = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    jpeg = io.BytesIO()
    Image.new("RGB", (300, 100), "white").save(jpeg, "JPEG")
    for image, text in [(empty, "пустая"), ("data:image/png;base64," + "A" * 200, "PNG"),
                        ("data:image/jpeg;base64," + base64.b64encode(jpeg.getvalue()).decode(), "PNG")]:
        resp = await client.put(f"/api/report/{rid}", headers=h,
                                json={"customer_signature": {"image": image, "signer_name": "Иванов"}})
        assert resp.status_code == 400 and text in resp.json()["detail"], resp.text

    resp = await client.patch(f"/api/report/{rid}/status", headers=h, json={"status_id": seeded["approved_id"]})
    assert resp.status_code == 200, resp.text
    resp = await client.put(f"/api/report/{rid}", headers=h, json={"customer_signature": SIG()})
    assert resp.status_code == 400 and "утверждён" in resp.json()["detail"]


async def test_act_context_has_signature(
    client: AsyncClient, db_session, superadmin_token, superadmin_user, auth_headers, reference_data,
):
    h = auth_headers(superadmin_token)
    seeded = await _seed(db_session, reference_data, superadmin_user["user"])
    resp = await client.post("/api/report/create", headers=h, json={
        "order_id": seeded["order_id"], "report_period": "2026-09",
        "customer_signature": SIG(signed_at="2026-09-18T11:32:00+03:00")})
    assert resp.status_code == 200, resp.text

    order = await render_docx._load_order_for_docx(db_session, seeded["order_id"])
    ctx = render_docx._build_context(order)
    assert ctx["report"]["is_signed"] is True
    assert ctx["report"]["signer_name"] == "Орлова Мария Сергеевна"
    assert ctx["report"]["signed_at"] == "18.09.2026 11:32"


async def test_rendered_act_contains_signature_image(
    client: AsyncClient, db_session, superadmin_token, superadmin_user, auth_headers, reference_data,
):
    """{{ customer_signature }} в шаблоне → картинка в готовом .docx."""
    import zipfile
    h = auth_headers(superadmin_token)
    user = superadmin_user["user"]
    seeded = await _seed(db_session, reference_data, user)
    resp = await client.post("/api/report/create", headers=h, json={
        "order_id": seeded["order_id"], "report_period": "2026-09", "customer_signature": SIG()})
    assert resp.status_code == 200, resp.text

    tpl = Document()
    tpl.add_paragraph("Заказчик: {{ report.signer_name }}, {{ report.signer_position }}")
    tpl.add_paragraph("{{ customer_signature }}")
    rel = "templates/test_signature_act.docx"
    (MEDIA_PATH / "templates").mkdir(parents=True, exist_ok=True)
    tpl.save(MEDIA_PATH / rel)
    order = await db_session.get(Order, seeded["order_id"])
    spec = await db_session.get(Spec_Order, order.spec_order_id)
    spec.template_storage_path = rel
    await db_session.commit()

    content, _name, _mt = await render_docx.render_order_document(seeded["order_id"], "docx", user)
    z = zipfile.ZipFile(io.BytesIO(content))
    assert any(n.startswith("word/media/") for n in z.namelist())
    text = " | ".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
    assert "Орлова Мария Сергеевна, Главный инженер" in text
